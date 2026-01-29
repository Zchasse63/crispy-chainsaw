#!/usr/bin/env python3
"""
Output Validation for Fleetwood Price Sheets

Run this BEFORE converting to PDF to catch errors early.
"""

import json
import os
import re
from typing import Dict, List, Tuple, Any
from openpyxl import load_workbook


def validate_pricing(items: List[Dict[str, Any]]) -> Tuple[List[str], List[str]]:
    """
    Validate pricing calculations.
    
    Args:
        items: List of item dictionaries with pricing data
    
    Returns:
        Tuple of (errors, warnings)
    """
    errors = []
    warnings = []
    
    for item in items:
        item_code = item.get('code', 'UNKNOWN')
        base_cost = item.get('base_cost', 0)
        landed = item.get('landed_per_lb')
        case_price = item.get('case_price')
        case_weight = item.get('case_weight')
        freight = item.get('freight_per_lb', 0)
        
        # ERROR: Margin must be applied
        if landed and base_cost and landed < (base_cost * 1.15):
            errors.append(f"{item_code}: Margin not applied correctly (landed ${landed:.2f} < base ${base_cost:.2f} × 1.15)")
        
        # WARNING: Freight outside typical range
        if freight and (freight < 0.10 or freight > 0.50):
            warnings.append(f"{item_code}: Freight ${freight:.2f}/lb outside typical range (0.10-0.50)")
        
        # WARNING: Case price outside typical range
        if case_price and isinstance(case_price, (int, float)) and (case_price < 5.00 or case_price > 500.00):
            warnings.append(f"{item_code}: Case price ${case_price:.2f} outside typical range (5.00-500.00)")
        
        # ERROR: Case math check
        if landed and case_weight and case_price and isinstance(case_price, (int, float)):
            expected_case = landed * case_weight
            variance = abs(case_price - expected_case) / expected_case if expected_case > 0 else 0
            if variance > 0.01:  # >1% variance
                errors.append(f"{item_code}: Case price ${case_price:.2f} doesn't match landed × weight (expected ${expected_case:.2f})")
    
    return errors, warnings


def validate_excel_file(excel_path: str, pallet_counts: Dict[str, Any] = None) -> Tuple[List[str], List[str], Dict[str, Any]]:
    """
    Validate data completeness in generated Excel file.
    
    Args:
        excel_path: Path to the Excel file
        pallet_counts: Optional pallet counts reference dict
    
    Returns:
        Tuple of (errors, warnings, stats)
    """
    errors = []
    warnings = []
    
    wb = load_workbook(excel_path)
    ws = wb.active
    
    total_items = 0
    items_with_cs_plt = 0
    items_with_nan = []
    internal_items_found = []
    cw_items_missing_marker = []
    items_with_price = 0
    
    # Skip header row
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        # Skip empty rows and warehouse headers
        if not row[0] or str(row[0]).upper() in ['AMERICUS, GA', 'BOYERTOWN, PA', 'INDIANAPOLIS, IN']:
            continue
        
        # Skip if it looks like a header or disclaimer
        if str(row[0]).upper() == 'CODE' or 'Prices include' in str(row[1] or ''):
            continue
        
        total_items += 1
        pack = str(row[2]).upper() if row[2] else ''
        desc = str(row[1]).upper() if row[1] else ''
        case_price = str(row[6]) if row[6] else ''
        per_lb = row[5] if len(row) > 5 else None
        
        # Check for NaN values
        for cell in row:
            cell_str = str(cell).lower() if cell else ''
            if 'nan' in cell_str or '#n/a' in cell_str:
                items_with_nan.append(row[0])
                break
        
        # Check CS/PLT populated (column 8, index 7)
        if len(row) > 7 and row[7] and str(row[7]).strip():
            items_with_cs_plt += 1
        
        # Check for internal items that should be excluded
        if 'TOTE' in pack or 'TOTE' in desc or 'TT' in pack:
            internal_items_found.append(row[0])
        
        # Check for large weight items (>=1000 LB)
        weight_match = re.search(r'(\d+)\s*LB', pack)
        if weight_match and int(weight_match.group(1)) >= 1000:
            internal_items_found.append(row[0])
        
        # Check CW markers for catch weight items
        cw_indicators = ['PC', '/HD', 'HEAD', 'TRAY']
        is_cw = any(ind in pack for ind in cw_indicators) or ' CW' in desc
        if is_cw and case_price.upper() != 'CW':
            cw_items_missing_marker.append(row[0])
        
        # Check price coverage
        if per_lb or (case_price and case_price.upper() != 'CW'):
            items_with_price += 1
    
    wb.close()
    
    # Evaluate results
    if items_with_nan:
        errors.append(f"NaN values found in {len(items_with_nan)} items: {items_with_nan[:3]}...")
    
    if internal_items_found:
        errors.append(f"Internal items not excluded ({len(internal_items_found)}): {internal_items_found[:3]}...")
    
    if cw_items_missing_marker:
        errors.append(f"CW markers missing ({len(cw_items_missing_marker)}): {cw_items_missing_marker[:3]}...")
    
    cs_plt_pct = (items_with_cs_plt / total_items * 100) if total_items > 0 else 0
    if cs_plt_pct == 0:
        errors.append("CS/PLT column is completely empty (0%)")
    elif cs_plt_pct < 50:
        warnings.append(f"CS/PLT coverage low: {items_with_cs_plt}/{total_items} ({cs_plt_pct:.0f}%)")
    
    price_pct = (items_with_price / total_items * 100) if total_items > 0 else 0
    if price_pct < 90:
        warnings.append(f"Price coverage low: {items_with_price}/{total_items} ({price_pct:.0f}%)")
    
    stats = {
        'total_items': total_items,
        'cs_plt_coverage': round(cs_plt_pct, 1),
        'price_coverage': round(price_pct, 1),
        'nan_count': len(items_with_nan),
        'internal_items': len(internal_items_found),
        'cw_missing': len(cw_items_missing_marker),
    }
    
    return errors, warnings, stats


def print_validation_report(errors: List[str], warnings: List[str], stats: Dict[str, Any]):
    """Print a formatted validation report."""
    print("\nVALIDATION RESULTS")
    print("=" * 50)
    print(f"Total items: {stats.get('total_items', 0)}")
    print(f"CS/PLT coverage: {stats.get('cs_plt_coverage', 0):.0f}%")
    print(f"Price coverage: {stats.get('price_coverage', 0):.0f}%")
    print(f"Errors: {len(errors)}")
    print(f"Warnings: {len(warnings)}")
    
    if errors:
        print(f"\n❌ ERRORS (must fix before PDF conversion):")
        for e in errors:
            print(f"  - {e}")
    
    if warnings:
        print(f"\n⚠️ WARNINGS (review recommended):")
        for w in warnings:
            print(f"  - {w}")
    
    if not errors and not warnings:
        print("\n✅ All validation checks passed!")
    elif not errors:
        print("\n✅ No errors - safe to proceed to PDF conversion")
    else:
        print("\n🛑 Fix errors before proceeding to PDF conversion")
    
    return len(errors) == 0


if __name__ == "__main__":
    import sys
    
    if len(sys.argv) < 2:
        print("Usage: python validate_output.py <excel_file>")
        sys.exit(1)
    
    excel_file = sys.argv[1]
    
    if not os.path.exists(excel_file):
        print(f"Error: File not found: {excel_file}")
        sys.exit(1)
    
    errors, warnings, stats = validate_excel_file(excel_file)
    success = print_validation_report(errors, warnings, stats)
    
    sys.exit(0 if success else 1)
